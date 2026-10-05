"""Generate the actionability table Excel (Korean and English) from config/actionability.yaml.

The YAML is the only source. Never edit the Excel files by hand; rerun this instead:
    python -m src.actionability_table
tests/test_actionability_excel.py fails if the committed Excel cells differ from the YAML.
"""
from __future__ import annotations

from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
YAML_PATH = ROOT / "config" / "actionability.yaml"
OUT = {lang: ROOT / "reports" / f"actionability_table_{lang}.xlsx" for lang in ("ko", "en")}

# Light tints per class (the class code is always written in the row too, so color is not the only cue)
CLASS_FILL = {
    "ACTIONABLE": "E3F4E8", "CONDITIONAL": "FFF4D6", "TIME_ONLY": "E3EEFB",
    "NON_DECREASING": "FBE6DE", "IMMUTABLE": "ECEBE8", "NOT_RECOMMENDED": "F3E3EE",
}
T = {
    "en": {
        "title": "Variable Actionability Classification",
        "generated": "Generated from config/actionability.yaml (version {v}, {d}). Do not edit by hand; "
                     "rerun `python -m src.actionability_table`. Educational use only.",
        "sheet_features": "Classification", "sheet_classes": "Classes",
        "headers": ["Variable", "Label", "Source", "Actionability", "Class", "Direction", "Step unit",
                    "Allowed range", "Difficulty (1-3)", "Months per step", "Monotone constraint (on PD)",
                    "Protected", "DiCE", "Recourse handling", "Notes"],
        "class_headers": ["Code", "Class", "Definition", "Recourse (DiCE) handling"],
        "source": {"gmsc": "GMSC", "derived": "Derived (Step 1)", "simulated": "Simulated"},
        "direction": {"increase": "Increase", "decrease": "Decrease", "none": "-"},
        "monotone": {1: "+1 (higher value, higher PD)", -1: "-1 (higher value, lower PD)", 0: "0 (none)"},
        "yes": "Yes", "vary": "Varied", "fixed": "Fixed", "no_upper": "no upper limit",
    },
    "ko": {
        "title": "변수별 조치 가능성 분류",
        "generated": "config/actionability.yaml에서 생성 (버전 {v}, {d}). 직접 수정하지 말고 "
                     "`python -m src.actionability_table`을 다시 실행하십시오. 교육용으로만 사용.",
        "sheet_features": "분류표", "sheet_classes": "분류 설명",
        "headers": ["변수", "설명", "출처", "조치 가능성", "분류명", "방향", "단위 변화량", "허용 범위",
                    "난이도 (1~3)", "단계당 개월", "단조 제약 (PD 기준)", "보호 속성", "DiCE", "경로 처리", "비고"],
        "class_headers": ["코드", "분류명", "정의", "경로(DiCE) 처리"],
        "source": {"gmsc": "GMSC", "derived": "파생 (1단계)", "simulated": "시뮬레이션"},
        "direction": {"increase": "증가", "decrease": "감소", "none": "-"},
        "monotone": {1: "+1 (값이 높을수록 PD 높음)", -1: "-1 (값이 높을수록 PD 낮음)", 0: "0 (없음)"},
        "yes": "예", "vary": "변경", "fixed": "고정", "no_upper": "상한 없음",
    },
}


def load_yaml(path: Path = YAML_PATH) -> dict:
    with open(path, encoding="utf-8") as f:
        return yaml.safe_load(f)


def _dash(v):
    return "-" if v is None else v


def _bounds(b, t) -> str:
    if b is None:
        return "-"
    return f"[{b[0]}, {t['no_upper'] if b[1] is None else b[1]}]"


def sheet_rows(act: dict, lang: str) -> dict[str, list[list]]:
    """All cell values per sheet, in order. Used by both the writer and the drift test."""
    t = T[lang]
    sfx = "ko" if lang == "ko" else "en"
    classes = act["classes"]
    header_rows = [[t["title"]], [t["generated"].format(v=act["version"], d=act["date"])], []]

    feats = [t["headers"]]
    for name, e in act["features"].items():
        cls = classes[e["actionability"]]
        feats.append([
            name, e[f"label_{sfx}"], t["source"][e["source"]], e["actionability"], cls[f"label_{sfx}"],
            t["direction"][e["direction"]], _dash(e["step"]), _bounds(e["bounds"], t), _dash(e["difficulty"]),
            _dash(e["months_per_step"]), t["monotone"][e["monotone"]], t["yes"] if e["protected"] else "-",
            t["vary"] if e["dice_vary"] else t["fixed"],
            cls["recourse"] if lang == "ko" else cls["recourse_en"], e[f"note_{sfx}"],
        ])
    cls_rows = [t["class_headers"]] + [
        [code, c[f"label_{sfx}"], c[f"definition_{sfx}"], c["recourse"] if lang == "ko" else c["recourse_en"]]
        for code, c in classes.items()
    ]
    return {t["sheet_features"]: header_rows + feats, t["sheet_classes"]: header_rows + cls_rows}


def write_workbook(act: dict, lang: str, path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    widths = {
        0: [38, 34, 16, 18, 20, 10, 12, 18, 12, 12, 26, 10, 10, 34, 90],
        1: [18, 22, 60, 50],
    }
    for i, (title, rows) in enumerate(sheet_rows(act, lang).items()):
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(r)
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"].font = Font(italic=True, color="52514E")
        header_row = 4
        for cell in ws[header_row]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9D8D2")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.freeze_panes = ws.cell(row=header_row + 1, column=2)
        for col, w in enumerate(widths[i], start=1):
            ws.column_dimensions[ws.cell(row=header_row, column=col).column_letter].width = w
        class_col = 4 if i == 0 else 1  # column holding the class code
        for row in ws.iter_rows(min_row=header_row + 1):
            fill = CLASS_FILL.get(row[class_col - 1].value)
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if fill:
                    cell.fill = PatternFill("solid", fgColor=fill)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> None:
    act = load_yaml()
    for lang, path in OUT.items():
        write_workbook(act, lang, path)
        print(f"Saved {path.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
