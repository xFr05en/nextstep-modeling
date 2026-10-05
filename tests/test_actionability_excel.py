"""The committed Excel files must match actionability.yaml cell by cell.

Compares cell values, not file bytes, because .xlsx files store timestamps.
"""
import pytest
from openpyxl import load_workbook

from src.actionability_table import OUT, load_yaml, sheet_rows


def _norm(row):
    vals = list(row)
    while vals and vals[-1] in (None, ""):
        vals.pop()
    return [None if v == "" else v for v in vals]


@pytest.mark.parametrize("lang", ["ko", "en"])
def test_excel_matches_yaml(lang):
    assert OUT[lang].exists(), f"{OUT[lang].name} missing (it is committed); rerun python -m src.actionability_table"
    wb = load_workbook(OUT[lang], read_only=True)
    expected = sheet_rows(load_yaml(), lang)
    assert wb.sheetnames == list(expected)
    for sheet, rows in expected.items():
        actual = [_norm(r) for r in wb[sheet].iter_rows(values_only=True)]
        exp = [_norm(r) for r in rows]
        assert len(actual) == len(exp), f"{lang}/{sheet}: row count differs, rerun python -m src.actionability_table"
        for i, (a, e) in enumerate(zip(actual, exp), start=1):
            assert a == e, f"{lang}/{sheet} row {i} differs from YAML, rerun python -m src.actionability_table"
